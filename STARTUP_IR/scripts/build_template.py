from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
P=Presentation();P.slide_width=Inches(13.333);P.slide_height=Inches(7.5);L=P.slide_layouts[6]
D='0A372D';G='11845D';M='E5F4EC';W='FFFFFF';S='F6FAF8';T='20332C';A='FFF1D8';Y='DEA53D';Q='5C716A';B='D4E6DC';FONT='Pretendard'
def col(h):return RGBColor.from_string(h)
def rect(s,x,y,w,h,c=W,b=None):
 a=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));a.fill.solid();a.fill.fore_color.rgb=col(c);a.line.color.rgb=col(b) if b else col(c);a.line.width=Pt(.8);return a
def tx(s,x,y,w,h,v,z=16,c=T,b=False):
 a=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=a.text_frame;f.clear();f.word_wrap=True
 for i,line in enumerate(v.split('\n')):
  p=f.paragraphs[0] if i==0 else f.add_paragraph();p.text=line;p.font.name=FONT;p.font.size=Pt(z);p.font.bold=b;p.font.color.rgb=col(c)
 return a
def line(s,x,y,x2,y2,c=G):
 a=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(x2),Inches(y2));a.line.color.rgb=col(c);a.line.width=Pt(1.4)
def circle(s,x,y,d,c=G):
 a=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));a.fill.solid();a.fill.fore_color.rgb=col(c);a.line.fill.background()
def base(n,title,sub):
 s=P.slides.add_slide(L);rect(s,0,0,13.333,7.5,W);rect(s,.34,.28,.08,.48,G);tx(s,.55,.26,11.8,.6,f'{n:02d}  |  {title}',26,T,True);tx(s,.55,.91,12,.38,sub,12,Q);tx(s,11.72,.12,1.3,.25,'GREEN IR',10,G,True);line(s,.38,7.1,12.9,7.1,B);tx(s,.4,7.14,11,.22,'[회사명] · 내부검토용 · 실제 성과/계획 구분 · 출처 기재',8,Q);tx(s,12.4,7.13,.5,.23,f'{n:02d}',9,Q);return s
def box(s,x,y,w,h,head,body='',highlight=False):
 rect(s,x,y,w,h,A if highlight else S,B);tx(s,x+.19,y+.2,w-.35,.46,head,17,D,True);tx(s,x+.2,y+.85,w-.4,h-.96,body,13,Q)
def arr(s,x,y):tx(s,x,y,.4,.45,'→',23,G,True)
# Cover
s=P.slides.add_slide(L);rect(s,0,0,13.333,7.5,S);rect(s,8.1,0,5.25,7.5,D)
for x,y,d,c in [(9.1,1.1,2.9,G),(10.6,2.9,2.2,'196F50'),(8.7,4.6,1.7,'4B9D78')]:circle(s,x,y,d,c)
tx(s,.8,1.2,6.7,.5,'[회사 로고 / COMPANY NAME]',19,G,True);tx(s,.8,2.15,6.7,1.7,'[한 문장으로 정의하는\n핵심 가치 제안]',34,D,True);tx(s,.8,4.55,6.5,.7,'[제품·기술 정의] · [핵심 고객]',18,Q);tx(s,.8,6.1,6.6,.45,'[대표자]  |  [작성일]  |  [연락처]',13,Q);tx(s,8.7,6.65,4,.4,'INVESTOR PRESENTATION',15,W,True)
# Problem
s=base(2,'PROBLEM','고객 과업 → 기존 대안 → 병목 → 사업적 손실')
for i,(h,b) in enumerate([('고객 과업','[고객의 핵심 목표]'),('기존 방식','[현행 프로세스]'),('핵심 병목','[실패·지연·비용]'),('사업적 손실','[고객 영향]')]):
 x=.6+i*3.14;box(s,x,2.05,2.75,2.4,h,b,i==2)
 if i<3:arr(s,x+2.79,3.03)
rect(s,.6,5.15,12.14,1.1,D);tx(s,.94,5.46,11.4,.55,'[고객이 구매할 이유를 한 문장으로 제시]',22,W,True)
# Solution
s=base(3,'SOLUTION','기존 방식 대비 고객 가치 · 기술 검증 범위 구분')
box(s,.6,1.65,3.15,4.75,'기존 방식','[기존 1]\n\n[기존 2]\n\n[구조적 한계]');tx(s,3.91,3.5,.6,.5,'VS',23,G,True)
rect(s,4.58,1.65,8.14,4.75,M,B);tx(s,4.95,1.93,7.5,.5,'[차별화된 솔루션]',20,D,True)
for i,h in enumerate(['요구조건','AI·기술 처리','검증','고객 산출물']):
 x=4.85+i*1.93;box(s,x,3.05,1.67,1.75,h,'[핵심 기능]')
 if i<3:arr(s,x+1.66,3.66)
tx(s,5.05,5.46,7.3,.35,'[검증된 차별점 / 개발 중인 차별점 구분]',14,G,True)
# Traction
s=base(4,'TRACTION','완료 · 진행 중 · 목표를 구분하여 검증 수치 중심 제시')
for i,(h,b) in enumerate([('고객·수요','[고객 수 / VOC]'),('PoC·성능','[측정 결과 / 기간]'),('유료 전환','[계약 / 매출]')]):
 x=.65+i*4.17;box(s,x,1.8,3.75,3.5,h,b);tx(s,x+.43,3.15,3.2,.6,'[핵심 KPI]',26,G,True)
rect(s,.65,5.65,12.1,.82,D);tx(s,.9,5.88,11.4,.4,'[가장 설득력 있는 실제 검증 성과]',18,W,True)
# Market
s=base(5,'MARKET SIZE','TAM / SAM / SOM 정의·지역·연도·산식·출처 필수')
for x,y,d,c in [(1.1,1.6,4.7,M),(1.8,2.3,3.3,'A9DFC4'),(2.55,3.05,1.8,G)]:circle(s,x,y,d,c)
tx(s,2.05,1.98,2.8,.4,'TAM  [금액]',17,D,True);tx(s,2.3,2.78,2.4,.4,'SAM  [금액]',17,D,True);tx(s,2.85,3.7,1.35,.45,'SOM',20,W,True)
for i,(h,b) in enumerate([('TAM','[전체 잠재시장 / 범위]'),('SAM','[서비스 가능시장 / 범위]'),('SOM','[고객수 × 도입률 × 단가]')]):box(s,6.4,1.65+i*1.5,6.2,1.27,h,b)
tx(s,6.55,6.4,5.9,.25,'출처: [기관·보고서·발표연도·페이지]',10,Q)
# BM
s=base(6,'BUSINESS MODEL','고객·제품·전달·수익의 현금흐름 구조')
for i,(h,b) in enumerate([('고객','[세그먼트]'),('가치 제공','[제품·서비스]'),('전달','[채널·파트너]'),('수익','[단가·주기]')]):
 x=.55+i*3.15;box(s,x,1.82,2.78,2.4,h,b)
 if i<3:arr(s,x+2.82,2.85)
for i,(h,b) in enumerate([('프로젝트형','[유료 PoC]'),('반복매출','[라이선스 / API]'),('고부가가치','[맞춤형 개발]')]):box(s,.55+i*4.17,4.78,3.78,1.45,h,b)
# GTM
s=base(7,'GO-TO-MARKET','유입 → PoC → 계약 → 확장, 단계별 전환 지표')
for i,(h,b) in enumerate([('유입','[접점]'),('검증','[PoC]'),('유료계약','[구매]'),('확장','[반복/업셀]')]):
 x=.55+i*3.15;box(s,x,2.05,2.78,2.65,h,b)
 if i<3:arr(s,x+2.82,3.12)
rect(s,.6,5.3,12.1,.9,D);tx(s,.87,5.58,11.5,.4,'[채널별 리드 · 전환율 · 고객획득비용 · 계약 주기]',16,W,True)
# Competition
s=base(8,'COMPETITION','비교 항목은 구매기준 · 미확인 기능을 ‘미지원’으로 단정 금지')
headers=['고객 구매기준','기존 방식','경쟁 대안','[자사]'];xs=[.6,4.0,7.13,10.2];ws=[3.25,3.0,2.93,2.52]
for x,w,h in zip(xs,ws,headers):rect(s,x,1.74,w,.65,D);tx(s,x+.13,1.9,w-.2,.35,h,14,W,True)
for i in range(4):
 y=2.5+i*.87
 for j,(x,w) in enumerate(zip(xs,ws)):
  rect(s,x,y,w,.78,M if j==3 else S,B);tx(s,x+.12,y+.2,w-.2,.4,f'[기준 {i+1}]' if j==0 else '[증빙]',13,D if j==3 else Q,j==3)
tx(s,.65,6.4,11.8,.3,'[비교 자료 출처 및 확인일]',10,Q)
# Team
s=base(9,'TEAM','핵심 실행역량 · 기술검증 파트너 · 미충원 역량')
for i,(h,b) in enumerate([('대표 / 사업','[시장·고객 이력]'),('기술 책임','[핵심 기술·특허]'),('사업화 책임','[제품·매출 이력]')]):
 x=.6+i*4.16;box(s,x,1.85,3.77,3.55,h,b);circle(s,x+1.42,3.0,.85,M)
rect(s,.6,5.75,12.2,.75,D);tx(s,.9,5.94,11.3,.38,'[검증 파트너 및 협력 관계: 계약·협의 상태 구분]',16,W,True)
# Financial
s=base(10,'FINANCIAL PLAN','5개년 매출·영업이익 · 가정과 산식 · 현금흐름')
box(s,.6,1.62,8.1,4.96,'5개년 재무 전망','[금액 단위]  ·  막대=매출  ·  선=영업이익')
for i in range(5):
 x=1.25+i*1.42;h=.5+i*.43;rect(s,x,5.52-h,.63,h,G);tx(s,x-.08,5.7,.9,.25,f'[Y{i+1}]',11,Q)
line(s,1.5,5.32,2.95,5.08);line(s,2.95,5.08,4.37,4.8);line(s,4.37,4.8,5.79,4.45);line(s,5.79,4.45,7.2,4.15)
box(s,8.94,1.62,3.8,2.25,'영업이익 산식','매출 − 매출원가 − 판관비')
box(s,8.94,4.05,3.8,2.53,'핵심 가정','[고객수 × 계약단가]\n[인건비·클라우드비]\n[고정비·변동비]')
# Scale
s=base(11,'SCALEUP / EXIT','마일스톤 · 자금조달 · 투자 회수 가능성')
for i,(h,b) in enumerate([('기술 검증','[성능·인증]'),('상업화','[유료 고객]'),('확장','[시장·파트너]'),('회수 경로','[M&A / IPO 가설]')]):
 x=.55+i*3.15;box(s,x,2.0,2.78,3.0,h,b)
 if i<3:arr(s,x+2.8,3.22)
rect(s,.6,5.48,12.1,.8,D);tx(s,.88,5.69,11.5,.42,'[각 단계의 정량 KPI·시점·필요 자금]',16,W,True)
# Funds
s=base(12,'USE OF FUNDS','투자 규모 · 배분 · 집행 시점 · 마일스톤 · 런웨이')
box(s,.6,1.65,3.2,4.8,'투자유치','[금액]\n\n[런웨이 개월]\n\n[조달 구조]')
for i,(h,b) in enumerate([('R&D·제품화','[금액 / %]'),('실증·인증','[금액 / %]'),('영업·시장','[금액 / %]'),('운영·인력','[금액 / %]')]):
 x=4.05+(i%2)*4.33;y=1.65+(i//2)*2.45;box(s,x,y,4.05,2.2,h,b)
tx(s,4.15,6.62,8.3,.25,'[집행 항목 합계 = 투자유치 금액]',11,Q)
OUT=Path(__file__).resolve().parents[1]/'templates'/'startup_ir_green_editable_12slides.pptx';P.save(OUT);print(OUT)
